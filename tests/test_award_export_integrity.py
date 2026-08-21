import io
import json
from decimal import Decimal
import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.main import app
from app import services
from award import exporter
from datasets.procurement_benchmark.dataset import get_benchmark_item_master


@pytest.fixture
def client():
    return TestClient(app)


def _setup_test_award_environment(rfq_id: str):
    """Sets up a clean RFQ with 2 suppliers and a finalized award."""
    award_file = services.get_award_file_path(rfq_id)
    if award_file.exists():
        award_file.unlink()

    item_master = get_benchmark_item_master()
    rfq_items = [
        {
            "sku": "BEAR-6205-2RS",
            "description": "Deep Groove Ball Bearing 25x52x15mm",
            "requested_quantity": "100",
            "requested_uom": "PCS",
            "target_price": "150.00"
        },
        {
            "sku": "CABL-6SQ-CU",
            "description": "Copper Armoured Flexible Cable 6 Sq.mm 4-Core",
            "requested_quantity": "500",
            "requested_uom": "MTR",
            "target_price": "200.00"
        }
    ]

    services.create_rfq(rfq_id, "Industrial Motors & Cables Project", "INR", rfq_items)

    # Supplier 1: Alpha Supplies
    q1 = services.create_quote("Alpha Supplies", rfq_id=rfq_id)
    csv1 = b"Line,SKU,Description,Qty,UOM,Unit Price\n1,BEAR-6205-2RS,Deep Groove Ball Bearing 25x52x15mm,100,PCS,120.00\n2,CABL-6SQ-CU,Copper Armoured Flexible Cable 6 Sq.mm 4-Core,500,MTR,190.00\n"
    services.upload_source_file(q1, "alpha.csv", csv1)
    services.run_extraction(q1)
    services.run_matching_for_quote(q1, rfq_id)
    services.resolve_match_candidate(q1, 0, "BEAR-6205-2RS", action="ACCEPT")
    services.resolve_match_candidate(q1, 1, "CABL-6SQ-CU", action="ACCEPT")

    # Supplier 2: Beta Industrial
    q2 = services.create_quote("Beta Industrial", rfq_id=rfq_id)
    csv2 = b"Line,SKU,Description,Qty,UOM,Unit Price\n1,BEAR-6205-2RS,Deep Groove Ball Bearing 25x52x15mm,100,PCS,115.00\n2,CABL-6SQ-CU,Copper Armoured Flexible Cable 6 Sq.mm 4-Core,500,MTR,195.00\n"
    services.upload_source_file(q2, "beta.csv", csv2)
    services.run_extraction(q2)
    services.run_matching_for_quote(q2, rfq_id)
    services.resolve_match_candidate(q2, 0, "BEAR-6205-2RS", action="ACCEPT")
    services.resolve_match_candidate(q2, 1, "CABL-6SQ-CU", action="ACCEPT")

    comp = services.run_rfq_comparison(rfq_id, [q1, q2], base_currency="INR")
    comp_id = comp["comparison_id"]

    # Finalize split award: Line 1 (Bearings) -> Beta (100 PCS @ 115), Line 2 (Cables) -> Alpha (500 MTR @ 190)
    proposal = services.build_proposed_award_allocation(rfq_id, scenario="LINE_ITEM_OPTIMAL")
    award = services.save_award_decision(rfq_id, proposal, is_finalized=True)
    return rfq_id, comp_id, q1, q2, award


def test_export_01_excel_award_workbook_sheets_and_values():
    rfq_id = "RFQ-EXP-01"
    _, comp_id, q1, q2, award = _setup_test_award_environment(rfq_id)

    wb_bytes = exporter.generate_award_excel_workbook(rfq_id)
    assert len(wb_bytes) > 1000

    wb = openpyxl.load_workbook(io.BytesIO(wb_bytes), data_only=True)
    assert "Executive Summary" in wb.sheetnames
    assert "Award Allocation" in wb.sheetnames
    assert "Supplier Summary" in wb.sheetnames
    assert "Audit & Provenance" in wb.sheetnames

    # Sheet 1: Executive Summary
    ws1 = wb["Executive Summary"]
    assert ws1["B5"].value == rfq_id
    assert "FINALIZED" in str(ws1["B7"].value)

    # Sheet 2: Award Allocation
    ws2 = wb["Award Allocation"]
    rows = list(ws2.iter_rows(values_only=True))
    assert rows[0][0] == "RFQ Line ID"
    assert len(rows) >= 3 # header + 2 lines
    # Verify values
    total_val = sum(Decimal(str(r[9])) for r in rows[1:])
    assert total_val == Decimal(award["total_awarded_value"])

    # Sheet 3: Supplier Summary
    ws3 = wb["Supplier Summary"]
    s_rows = list(ws3.iter_rows(values_only=True))
    assert len(s_rows) == 3 # header + 2 suppliers (Alpha & Beta)
    suppliers_found = {r[0] for r in s_rows[1:]}
    assert "Alpha Supplies" in suppliers_found
    assert "Beta Industrial" in suppliers_found

    # Sheet 4: Audit & Provenance
    ws4 = wb["Audit & Provenance"]
    assert ws4["B4"].value == rfq_id
    assert ws4["B6"].value == comp_id


def test_export_02_pdf_executive_award_report_generation():
    rfq_id = "RFQ-EXP-02"
    _setup_test_award_environment(rfq_id)

    pdf_bytes = exporter.generate_award_executive_pdf(rfq_id)
    assert pdf_bytes.startswith(b"%PDF-")
    assert len(pdf_bytes) > 2000


def test_export_03_supplier_specific_requisitions_excel_and_pdf():
    rfq_id = "RFQ-EXP-03"
    _, _, q1, q2, _ = _setup_test_award_environment(rfq_id)

    # Excel requisition for Alpha
    req_xlsx = exporter.generate_supplier_award_requisition_excel(rfq_id, q1)
    wb = openpyxl.load_workbook(io.BytesIO(req_xlsx), data_only=True)
    ws = wb["Award Requisition"]
    assert "PURCHASE ORDER REQUISITION" in ws["A1"].value
    assert "NOT A LEGALLY BINDING PURCHASE ORDER" in ws["A2"].value

    # PDF requisition for Beta
    req_pdf = exporter.generate_supplier_award_requisition_pdf(rfq_id, q2)
    assert req_pdf.startswith(b"%PDF-")
    assert len(req_pdf) > 1000


def test_export_04_csv_export_schema_and_integrity():
    rfq_id = "RFQ-EXP-04"
    _, comp_id, q1, q2, award = _setup_test_award_environment(rfq_id)

    csv_str = exporter.generate_award_csv_export(rfq_id)
    lines = csv_str.strip().split("\n")
    header = lines[0].split(",")
    assert "rfq_id" in header
    assert "unit_landed_cost" in header
    assert "line_total_value" in header

    assert len(lines) == 3 # header + 2 allocated lines
    # Grand total check
    csv_total = sum(Decimal(l.split(",")[11]) for l in lines[1:])
    assert csv_total == Decimal(award["total_awarded_value"])


def test_export_05_reopened_and_refinalized_exports_latest_state():
    rfq_id = "RFQ-EXP-05"
    _, comp_id, q1, q2, award1 = _setup_test_award_environment(rfq_id)
    total_val_1 = Decimal(award1["total_awarded_value"])

    # Reopen
    services.reopen_award_decision(rfq_id, reason="Modifying quantities for project scope reduction")

    # Modify line 1 quantity: Beta only gets 50 PCS instead of 100 PCS (partial line)
    proposal = services.get_award_decision(rfq_id)
    proposal["allocations"][0]["supplier_splits"][0]["allocated_qty"] = "50"
    proposal["buyer_accepted_unallocated"] = True

    award2 = services.save_award_decision(rfq_id, proposal, is_finalized=True)
    total_val_2 = Decimal(award2["total_awarded_value"])
    assert total_val_2 < total_val_1

    # Export workbook and verify updated total and revision history
    wb_bytes = exporter.generate_award_excel_workbook(rfq_id)
    wb = openpyxl.load_workbook(io.BytesIO(wb_bytes), data_only=True)
    ws4 = wb["Audit & Provenance"]
    
    # Revision log should show revision 1 with reason
    r_rows = list(ws4.iter_rows(values_only=True))
    assert any("Modifying quantities for project scope reduction" in str(cell) for row in r_rows for cell in row)


def test_export_06_http_routes(client: TestClient):
    rfq_id = "RFQ-EXP-06"
    _, comp_id, q1, q2, award = _setup_test_award_environment(rfq_id)

    # 1. Excel export endpoint
    res_xlsx = client.get(f"/rfqs/{rfq_id}/award/export/excel")
    assert res_xlsx.status_code == 200
    assert res_xlsx.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert f"{rfq_id}_Award_Workbook.xlsx" in res_xlsx.headers["content-disposition"]

    # 2. PDF export endpoint
    res_pdf = client.get(f"/rfqs/{rfq_id}/award/export/pdf")
    assert res_pdf.status_code == 200
    assert res_pdf.headers["content-type"] == "application/pdf"
    assert f"{rfq_id}_Executive_Award_Report.pdf" in res_pdf.headers["content-disposition"]

    # 3. CSV export endpoint
    res_csv = client.get(f"/rfqs/{rfq_id}/award/export/csv")
    assert res_csv.status_code == 200
    assert res_csv.headers["content-type"] == "text/csv; charset=utf-8"
    assert f"{rfq_id}_Award_Allocation.csv" in res_csv.headers["content-disposition"]

    # 4. Supplier Requisition Excel & PDF endpoints
    res_req_xlsx = client.get(f"/rfqs/{rfq_id}/award/export/requisition/{q1}/excel")
    assert res_req_xlsx.status_code == 200

    res_req_pdf = client.get(f"/rfqs/{rfq_id}/award/export/requisition/{q1}/pdf")
    assert res_req_pdf.status_code == 200


def test_export_07_cross_rfq_isolation():
    rfq1 = "RFQ-EXP-ISOL-01"
    rfq2 = "RFQ-EXP-ISOL-02"
    _setup_test_award_environment(rfq1)
    _setup_test_award_environment(rfq2)

    csv1 = exporter.generate_award_csv_export(rfq1)
    csv2 = exporter.generate_award_csv_export(rfq2)

    assert rfq1 in csv1
    assert rfq2 not in csv1
    assert rfq2 in csv2
    assert rfq1 not in csv2
