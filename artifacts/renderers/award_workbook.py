"""
Domain-Agnostic Multi-Tab Excel Award Workbook Renderer.
Consumes pre-computed AwardRecordViewModel. Performs NO business calculations.
"""

from datetime import datetime, timezone
from decimal import Decimal
import io
from typing import Tuple

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from artifacts.models import ArtifactFormat, ArtifactMetadata, ArtifactType
from artifacts.registry import BaseArtifactRenderer, register_renderer
from artifacts.view_models import AwardRecordViewModel


class AwardWorkbookRenderer(BaseArtifactRenderer):
    artifact_type = ArtifactType.AWARD_WORKBOOK

    def render(self, vm: AwardRecordViewModel, **kwargs) -> Tuple[bytes, ArtifactMetadata]:
        wb = openpyxl.Workbook()
        
        # Styles
        navy_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        accent_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
        subtle_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        
        white_bold = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        navy_bold = Font(name="Calibri", size=11, bold=True, color="0F172A")
        title_font = Font(name="Calibri", size=14, bold=True, color="0F172A")
        subtitle_font = Font(name="Calibri", size=10, italic=True, color="64748B")
        data_font = Font(name="Calibri", size=10, color="1E293B")
        bold_data_font = Font(name="Calibri", size=10, bold=True, color="0F172A")
        
        thin_border_side = Side(style="thin", color="CBD5E1")
        grid_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

        # ---------------------------------------------------------------------
        # Sheet 1: Executive Summary
        # ---------------------------------------------------------------------
        ws1 = wb.active
        ws1.title = "Executive Summary"
        ws1.views.sheetView[0].showGridLines = True
        
        ws1.append(["PROCUREMENT AWARD DECISION — EXECUTIVE SUMMARY"])
        ws1["A1"].font = title_font
        ws1.append([f"Authoritative evaluation snapshot &bull; Project: {vm.rfq_title}"])
        ws1["A2"].font = subtitle_font
        ws1.append([])

        summary_rows = [
            ("RFQ Reference Number", vm.rfq_id),
            ("RFQ Project Title", vm.rfq_title),
            ("Award Status", vm.status),
            ("Finalized Date", vm.finalized_at_display),
            ("Awarded By", vm.buyer_id),
            ("Base Currency", vm.base_currency),
            ("Total Award Commitment", vm.total_awarded_value_display),
            ("Total Unique RFQ Lines", vm.fulfillment.total_rfq_lines),
            ("Fully Fulfilled Lines", vm.fulfillment.fully_allocated_lines),
            ("Partially Fulfilled Lines", vm.fulfillment.partially_allocated_lines),
            ("Unallocated Scope Lines", vm.fulfillment.unallocated_lines),
            ("Overall Scope Fulfillment", vm.fulfillment.overall_fulfillment_display),
            ("Sourcing Strategy", vm.sourcing_strategy_label),
            ("Commercial Comparison ID", vm.comparison_id),
        ]

        ws1.append(["Parameter", "Value"])
        ws1["A4"].fill = navy_fill
        ws1["A4"].font = white_bold
        ws1["B4"].fill = navy_fill
        ws1["B4"].font = white_bold

        for row_idx, (k, v) in enumerate(summary_rows, start=5):
            ws1.append([k, str(v)])
            ws1[f"A{row_idx}"].font = bold_data_font
            ws1[f"B{row_idx}"].font = data_font
            ws1[f"A{row_idx}"].border = grid_border
            ws1[f"B{row_idx}"].border = grid_border
            if row_idx % 2 == 0:
                ws1[f"A{row_idx}"].fill = subtle_fill
                ws1[f"B{row_idx}"].fill = subtle_fill

        ws1.column_dimensions["A"].width = 30
        ws1.column_dimensions["B"].width = 50

        # ---------------------------------------------------------------------
        # Sheet 2: Award Allocation
        # ---------------------------------------------------------------------
        ws2 = wb.create_sheet(title="Award Allocation")
        ws2.views.sheetView[0].showGridLines = True
        
        alloc_headers = [
            "RFQ Line ID", "Item SKU", "Description & Specifications", "Required Qty", "UOM",
            "Allocation Type", "Awarded Supplier", "Quote Reference", "Split Awarded Qty", "Unit Landed Cost",
            "Split Award Value", "Currency", "Allocation State", "L1 Optimal"
        ]
        ws2.append(alloc_headers)
        for col_idx in range(1, len(alloc_headers) + 1):
            cell = ws2.cell(row=1, column=col_idx)
            cell.fill = navy_fill
            cell.font = white_bold
            cell.alignment = Alignment(horizontal="center", vertical="center")

        curr_row = 2
        for line in vm.lines:
            if line.splits:
                for sp in line.splits:
                    is_first = (sp.split_index == 1)
                    sku_disp = f"{line.item_sku} [{sp.allocation_type_label}]" if not sp.is_single else line.item_sku
                    req_disp = float(line.required_qty_display.replace(",", "")) if (is_first and line.required_qty_display != "N/A") else 0.0
                    
                    ws2.append([
                        line.line_id, sku_disp, line.item_description, req_disp, line.canonical_uom,
                        sp.allocation_type_label, sp.supplier_name, sp.quote_id,
                        float(sp.awarded_qty_display.replace(",", "")),
                        float(sp.unit_landed_cost_display.replace(",", "")),
                        float(sp.split_value_display.replace(",", "")),
                        vm.base_currency, line.allocation_state, "YES" if sp.is_l1 else "NO"
                    ])
                    
                    ws2.cell(row=curr_row, column=4).number_format = "#,##0.00"
                    ws2.cell(row=curr_row, column=9).number_format = "#,##0.00"
                    ws2.cell(row=curr_row, column=10).number_format = "#,##0.0000"
                    ws2.cell(row=curr_row, column=11).number_format = "#,##0.00"
                    for c in range(1, len(alloc_headers) + 1):
                        ws2.cell(row=curr_row, column=c).border = grid_border
                        ws2.cell(row=curr_row, column=c).font = data_font
                    curr_row += 1
            else:
                ws2.append([
                    line.line_id, line.item_sku, line.item_description,
                    float(line.required_qty_display.replace(",", "")) if line.required_qty_display != "N/A" else 0.0,
                    line.canonical_uom, "UNALLOCATED", "UNALLOCATED", "N/A", 0.0, 0.0, 0.0,
                    vm.base_currency, line.allocation_state, "NO"
                ])
                for c in range(1, len(alloc_headers) + 1):
                    ws2.cell(row=curr_row, column=c).border = grid_border
                    ws2.cell(row=curr_row, column=c).font = data_font
                curr_row += 1

        for col in ws2.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws2.column_dimensions[col_letter].width = max(max_len + 3, 12)

        # ---------------------------------------------------------------------
        # Sheet 3: Supplier Summary (Dimension-Safe, Unique RFQ Lines)
        # ---------------------------------------------------------------------
        ws3 = wb.create_sheet(title="Supplier Summary")
        ws3.views.sheetView[0].showGridLines = True
        
        supp_headers = ["Supplier Name", "Quote ID", "RFQ Lines Awarded", "Split Allocations", "Quantities by UOM", f"Total Award Value ({vm.base_currency})", "Percentage Share"]
        ws3.append(supp_headers)
        for col_idx in range(1, len(supp_headers) + 1):
            cell = ws3.cell(row=1, column=col_idx)
            cell.fill = navy_fill
            cell.font = white_bold
            cell.alignment = Alignment(horizontal="center", vertical="center")

        s_row = 2
        for s in vm.suppliers:
            ws3.append([
                s.supplier_name, s.quote_id, s.rfq_lines_awarded_count, s.split_allocations_count,
                s.quantities_by_uom_display.replace("&bull;", "•"),
                float(s.total_value_display.replace(",", "")),
                s.percentage_share_display
            ])
            ws3.cell(row=s_row, column=6).number_format = "#,##0.00"
            for c in range(1, len(supp_headers) + 1):
                ws3.cell(row=s_row, column=c).border = grid_border
                ws3.cell(row=s_row, column=c).font = data_font
            s_row += 1

        for col in ws3.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws3.column_dimensions[col_letter].width = max(max_len + 3, 14)

        buf = io.BytesIO()
        wb.save(buf)
        excel_bytes = buf.getvalue()

        meta = ArtifactMetadata(
            artifact_id=f"ART-XLS-{vm.award_id}",
            award_id=vm.award_id,
            rfq_id=vm.rfq_id,
            artifact_type=ArtifactType.AWARD_WORKBOOK,
            format=ArtifactFormat.EXCEL,
            name="Award Workbook",
            filename=f"{vm.rfq_id}_Award_Workbook.xlsx",
            mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            size_bytes=len(excel_bytes),
            download_url=f"/rfqs/{vm.rfq_id}/award/export/excel"
        )

        return excel_bytes, meta


register_renderer(ArtifactType.AWARD_WORKBOOK, AwardWorkbookRenderer)
