import io
import json
from pathlib import Path
from decimal import Decimal
import openpyxl

base_dir = Path(r"C:\Users\ABHIMANYU KUMAR\.gemini\antigravity\scratch\quote_intelligence")
syn_dir = base_dir / "datasets" / "synthetic"
gt_dir = base_dir / "datasets" / "ground_truth"

syn_dir.mkdir(parents=True, exist_ok=True)
gt_dir.mkdir(parents=True, exist_ok=True)

# -------------------------------------------------------------
# Fixture 1: 01_clean_quote.xlsx
# -------------------------------------------------------------
wb1 = openpyxl.Workbook()
ws1 = wb1.active
ws1.title = "Quotation"
ws1.append(["Premier Industrial Fasteners Ltd"])
ws1.append(["Quote Ref:", "QT-PIF-2026-001", "Date:", "2026-08-19"])
ws1.append([])
ws1.append(["Item Code", "Description", "Qty", "UOM", "Rate", "Tax %"])
ws1.append(["BOLT-M8", "Hex Bolt M8x25 SS304", 100, "PCS", 4.50, 18.0])
ws1.append(["NUT-M8", "Hex Nut M8 SS304", 100, "NOS", 2.00, 18.0])
wb1.save(syn_dir / "01_clean_quote.xlsx")

gt1 = {
    "quote_id": "GT-01",
    "supplier_raw_name": "Premier Industrial Fasteners Ltd",
    "quote_number": "QT-PIF-2026-001",
    "quote_date": "2026-08-19",
    "currency": "INR",
    "exchange_rate_to_base": "1.0",
    "items": [
        {
            "line_index": 0,
            "raw_description": "Hex Bolt M8x25 SS304",
            "supplier_part_number": "BOLT-M8",
            "quoted_qty": "100",
            "quoted_uom": "PCS",
            "unit_price": "4.5",
            "tax_rate_pct": "18"
        },
        {
            "line_index": 1,
            "raw_description": "Hex Nut M8 SS304",
            "supplier_part_number": "NUT-M8",
            "quoted_qty": "100",
            "quoted_uom": "PCS",
            "unit_price": "2.0",
            "tax_rate_pct": "18"
        }
    ],
    "extraction_metadata": {
        "source_file_name": "01_clean_quote.xlsx",
        "source_file_hash": "gt_hash_01",
        "parser_used": "ground_truth"
    }
}
with open(gt_dir / "01_clean_quote.json", "w", encoding="utf-8") as f:
    json.dump(gt1, f, indent=2)

# -------------------------------------------------------------
# Fixture 2: 02_multi_sheet_quote.xlsx (Sheet 1 = Cover, Sheet 2 = Quotation)
# -------------------------------------------------------------
wb2 = openpyxl.Workbook()
ws2_cover = wb2.active
ws2_cover.title = "Company Info"
ws2_cover.append(["Zenith Machinery Spares Pvt Ltd"])
ws2_cover.append(["Registered Address: Plot 44, Peenya Industrial Area, Bangalore"])
ws2_cover.append(["General Terms: Payment within 30 days of invoice"])

ws2_quote = wb2.create_sheet(title="Quote Details")
ws2_quote.append(["Zenith Machinery Spares Pvt Ltd"])
ws2_quote.append(["Quotation No:", "QT-ZMS-2026-778", "Date:", "2026-08-19"])
ws2_quote.append([])
ws2_quote.append(["Description", "Part #", "Quantity", "Unit", "Unit Rate", "GST %"])
ws2_quote.append(["Heavy Duty Spur Gear 24T", "GEAR-24T", 5, "Nos", 2400.00, 18.0])
ws2_quote.append(["Flange Bushing 30x40", "BUSH-3040", 20, "Pieces", 320.00, 18.0])
wb2.save(syn_dir / "02_multi_sheet_quote.xlsx")

gt2 = {
    "quote_id": "GT-02",
    "supplier_raw_name": "Zenith Machinery Spares Pvt Ltd",
    "quote_number": "QT-ZMS-2026-778",
    "quote_date": "2026-08-19",
    "currency": "INR",
    "exchange_rate_to_base": "1.0",
    "items": [
        {
            "line_index": 0,
            "raw_description": "Heavy Duty Spur Gear 24T",
            "supplier_part_number": "GEAR-24T",
            "quoted_qty": "5",
            "quoted_uom": "PCS",
            "unit_price": "2400",
            "tax_rate_pct": "18"
        },
        {
            "line_index": 1,
            "raw_description": "Flange Bushing 30x40",
            "supplier_part_number": "BUSH-3040",
            "quoted_qty": "20",
            "quoted_uom": "PCS",
            "unit_price": "320",
            "tax_rate_pct": "18"
        }
    ],
    "extraction_metadata": {
        "source_file_name": "02_multi_sheet_quote.xlsx",
        "source_file_hash": "gt_hash_02",
        "parser_used": "ground_truth"
    }
}
with open(gt_dir / "02_multi_sheet_quote.json", "w", encoding="utf-8") as f:
    json.dump(gt2, f, indent=2)

# -------------------------------------------------------------
# Fixture 3: 03_multiline_desc_quote.xlsx (Multi-line descriptions)
# -------------------------------------------------------------
wb3 = openpyxl.Workbook()
ws3 = wb3.active
ws3.title = "Quote"
ws3.append(["Delta Electricals India"])
ws3.append(["Quotation Ref:", "QT-DEL-554", "Dated:", "2026-08-19"])
ws3.append([])
ws3.append(["Description", "Qty", "UOM", "Rate", "Tax %"])
ws3.append(["3 Phase Induction Motor 10HP 1440 RPM", 2, "NOS", 18500.00, 18.0])
ws3.append(["IE3 High Efficiency Premium Model - Cast Iron Body", "", "", "", ""])  # Continuation line
ws3.append(["Direct On-Line (DOL) Motor Starter 10HP", 2, "NOS", 3200.00, 18.0])
wb3.save(syn_dir / "03_multiline_desc_quote.xlsx")

gt3 = {
    "quote_id": "GT-03",
    "supplier_raw_name": "Delta Electricals India",
    "quote_number": "QT-DEL-554",
    "quote_date": "2026-08-19",
    "currency": "INR",
    "exchange_rate_to_base": "1.0",
    "items": [
        {
            "line_index": 0,
            "raw_description": "3 Phase Induction Motor 10HP 1440 RPM IE3 High Efficiency Premium Model - Cast Iron Body",
            "quoted_qty": "2",
            "quoted_uom": "PCS",
            "unit_price": "18500",
            "tax_rate_pct": "18"
        },
        {
            "line_index": 1,
            "raw_description": "Direct On-Line (DOL) Motor Starter 10HP",
            "quoted_qty": "2",
            "quoted_uom": "PCS",
            "unit_price": "3200",
            "tax_rate_pct": "18"
        }
    ],
    "extraction_metadata": {
        "source_file_name": "03_multiline_desc_quote.xlsx",
        "source_file_hash": "gt_hash_03",
        "parser_used": "ground_truth"
    }
}
with open(gt_dir / "03_multiline_desc_quote.json", "w", encoding="utf-8") as f:
    json.dump(gt3, f, indent=2)

# -------------------------------------------------------------
# Fixture 4: 04_clean_quote.csv
# -------------------------------------------------------------
csv_text = """# Supplier: Matrix Bearings Corp
# Quote Number: QT-MAT-2026-303
# Date: 2026-08-19

Item Code,Description,Qty,UOM,Rate,Tax %
6205-2RS,"SKF Ball Bearing 6205 2RS C3",25,NOS,520.00,18%
6306-2RS,"SKF Ball Bearing 6306 2RS C3",15,NOS,780.00,18%
"""
with open(syn_dir / "04_clean_quote.csv", "w", encoding="utf-8") as f:
    f.write(csv_text)

gt4 = {
    "quote_id": "GT-04",
    "supplier_raw_name": "Matrix Bearings Corp",
    "quote_number": "QT-MAT-2026-303",
    "quote_date": "2026-08-19",
    "currency": "INR",
    "exchange_rate_to_base": "1.0",
    "items": [
        {
            "line_index": 0,
            "raw_description": "SKF Ball Bearing 6205 2RS C3",
            "supplier_part_number": "6205-2RS",
            "quoted_qty": "25",
            "quoted_uom": "PCS",
            "unit_price": "520",
            "tax_rate_pct": "18"
        },
        {
            "line_index": 1,
            "raw_description": "SKF Ball Bearing 6306 2RS C3",
            "supplier_part_number": "6306-2RS",
            "quoted_qty": "15",
            "quoted_uom": "PCS",
            "unit_price": "780",
            "tax_rate_pct": "18"
        }
    ],
    "extraction_metadata": {
        "source_file_name": "04_clean_quote.csv",
        "source_file_hash": "gt_hash_04",
        "parser_used": "ground_truth"
    }
}
with open(gt_dir / "04_clean_quote.json", "w", encoding="utf-8") as f:
    json.dump(gt4, f, indent=2)

print("Generated 4 synthetic benchmark files and golden JSON ground truths.")
